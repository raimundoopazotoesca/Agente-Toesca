"""Tests del pipeline JLL v2: parser, idempotencia, atomicidad y lineage.

Los fixtures construyen planillas .xlsx minimas en memoria en vez de depender
del archivo real de JLL, que no es la version oficial y vive fuera del repo.
"""
import sqlite3

import openpyxl
import pytest

from tools.db import repo_jll_v2, repo_rent_roll
from tools.db.connection import apply_migrations, get_conn_for
from tools.jll_planilla_tools import (
    FUENTE_FORMATO,
    FUENTE_PROVEEDOR,
    es_formato_v2,
    parse_planilla,
)

RR_HEADER = [
    "portafolio", "fecha_corte_rent_roll", "piso", "local", "categoría_general",
    "#Contrato", "estado_rent_roll", "información_locatario_gla",
    "información_locatario_marca", "información_locatario_razón_social",
    "información_locatario_estado_contrato", "fechas_de_ocupación_inicio",
    "fechas_de_ocupación_fin", "renta_renta_real", "renta_m2",
]
AUX_HEADER = [
    "portafolio", "mes", "rubro_presupuestal", "clasificación presupuestal",
    "código_contable", "nombre_del_tercero", "descripción", "débito", "crédito",
]
CART_HEADER = [
    "Portafolio ", "cliente", "marca", "identificación", "documento",
    "concepto", "inmueble", "fecha_vencimiento", "fecha_de_corte",
    "días_de_vencimiento", "vencido_1_a_30", "vencido_31_a_60",
    "vencido_61_a_90", "vencido_más_de_91", "saldo_por_vencer",
    "saldo_a_favor", "total_cartera",
]
REC_HEADER = ["Proyecto", "Fecha corte", "Valor recaudado"]


def _rr_row(portafolio="Torre A", corte="2026-07-01", local="101",
            categoria="OFICINA", estado="Ocupado", gla=100.0, total=60.0,
            tasa=0.6, razon="ACME SPA"):
    return [portafolio, corte, "PISO 1", local, categoria, "C1", estado, gla,
            razon, razon, "Firmado", "2020-01-01", "2027-01-31", total, tasa]


def _escribir(path, rr_rows, aux_rows=None, cart_rows=None, rec_rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RentRoll"
    ws.append(RR_HEADER)
    for r in rr_rows:
        ws.append(r)

    aux = wb.create_sheet("AuxiliarContable")
    aux.append(AUX_HEADER)
    for r in aux_rows or []:
        aux.append(r)

    cart = wb.create_sheet("Cartera")
    cart.append(CART_HEADER)
    for r in cart_rows or []:
        cart.append(r)

    rec = wb.create_sheet("Recaudado")
    rec.append(REC_HEADER)
    for r in rec_rows or []:
        rec.append(r)

    wb.save(path)
    return str(path)


@pytest.fixture
def planilla(tmp_path):
    return _escribir(
        tmp_path / "jll.xlsx",
        rr_rows=[_rr_row(), _rr_row(local="102", estado="Vacante", total=0, tasa=0)],
        aux_rows=[
            ["Torre A", "2026-07-31", "Ingreso por arriendo", "Ingresos", 1,
             "ACME SPA", "Ingreso por arriendo", None, 60.0],
            ["Torre A", "2026-07-31", "Ingreso por intereses", "Ingresos", 9,
             "ACME SPA", "Ingreso por intereses", None, 1.5],
        ],
        cart_rows=[
            ["Torre A", "ACME SPA", "ACME", "76158201", "SI 1", "RENTAS",
             "LOCAL 1", "2026-06-01", "2026-07-31", 30, 10.0, 0.0, 0.0, 0.0,
             0.0, 0.0, 10.0],
        ],
        rec_rows=[["Torre A", "2026-07-31", 55.0]],
    )


@pytest.fixture
def db(tmp_db_path):
    apply_migrations(tmp_db_path)
    conn = get_conn_for(tmp_db_path)
    conn.execute(
        "INSERT OR IGNORE INTO dim_activo (activo_key, nombre, fondo_key) "
        "VALUES ('Torre A', 'Torre A', 'PT')"
    )
    conn.commit()
    conn.close()
    return tmp_db_path


# ── parser ──────────────────────────────────────────────────────────────────

def test_detecta_formato_v2(planilla):
    assert es_formato_v2(planilla) is True


def test_parser_normaliza_periodo_y_conserva_fecha_fuente(planilla):
    p = parse_planilla(planilla)
    fila = p.rent_roll.filas[0]
    # Convencion acordada: dia 1 identifica el mes en curso.
    assert fila["periodo"] == "2026-07"
    # La fecha original se conserva para poder recorregir la interpretacion.
    assert fila["fecha_corte_fuente"] == "2026-07-01"


def test_parser_separa_total_y_tasa(planilla):
    fila = parse_planilla(planilla).rent_roll.filas[0]
    assert fila["renta_uf"] == 60.0        # total UF
    assert fila["renta_uf_m2"] == 0.6      # tasa UF/m2
    assert fila["renta_semantica"] == "total_uf"


def test_scope_sale_del_contenido(planilla):
    p = parse_planilla(planilla)
    assert p.scope() == {(FUENTE_PROVEEDOR, "Torre A", "2026-07")}


def test_estado_invalido_es_anomalia_y_no_descarte(tmp_path):
    path = _escribir(tmp_path / "malo.xlsx", [_rr_row(estado="BS4B07")])
    p = parse_planilla(path)
    tipos = {a.tipo for a in p.rent_roll.anomalias}
    assert "estado_invalido" in tipos
    # fail-loud, no silencioso: la fila igual se conserva
    assert len(p.rent_roll.filas) == 1


def test_portafolio_desconocido_es_anomalia(tmp_path):
    path = _escribir(tmp_path / "raro.xlsx", [_rr_row(portafolio="Edificio X")])
    p = parse_planilla(path)
    assert {a.tipo for a in p.rent_roll.anomalias} == {"portafolio_desconocido"}
    assert p.rent_roll.filas[0]["activo_key"] is None


def test_categoria_ug_queda_pendiente_no_asumida(tmp_path):
    path = _escribir(tmp_path / "ug.xlsx", [_rr_row(categoria="UG")])
    p = parse_planilla(path)
    assert "categoria_ug_pendiente" in {a.tipo for a in p.rent_roll.anomalias}


# ── validate ────────────────────────────────────────────────────────────────

def test_validate_no_escribe_nada(planilla, db):
    from tools.db.ingest_jll_planilla import validate

    validate(planilla, db)
    conn = get_conn_for(db)
    assert conn.execute("SELECT COUNT(*) FROM raw_rent_roll_line").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM ingest_run").fetchone()[0] == 0
    conn.close()


def test_validate_reporta_rubros_sin_mapping(planilla, db):
    from tools.db.ingest_jll_planilla import validate

    r = validate(planilla, db)
    assert any(
        "Ingreso por intereses" in x for x in r["mapeo_er"]["rubros_sin_mapping"]
    )


# ── idempotencia y retry ────────────────────────────────────────────────────

def test_commit_carga_las_cuatro_hojas(planilla, db):
    from tools.db.ingest_jll_planilla import commit

    r = commit(planilla, db)
    assert r["status"] == "ok"
    conn = get_conn_for(db)
    assert conn.execute("SELECT COUNT(*) FROM raw_rent_roll_line").fetchone()[0] == 2
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_movimiento_contable_line").fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM raw_cartera_line").fetchone()[0] == 1
    assert conn.execute("SELECT COUNT(*) FROM raw_recaudacion").fetchone()[0] == 1
    conn.close()


def test_mismo_hash_tras_exito_es_noop(planilla, db):
    """Caso 1 del contrato: duplicado -> NO-OP, sin supersede ni insert."""
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db)
    conn = get_conn_for(db)
    antes = conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE superseded_at IS NULL"
    ).fetchone()[0]
    conn.close()

    r2 = commit(planilla, db)
    assert r2["status"] == "skipped_duplicate"

    conn = get_conn_for(db)
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE superseded_at IS NULL"
    ).fetchone()[0] == antes
    # No-op de verdad: nada quedo superseded y no se abrio un run nuevo.
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE superseded_at IS NOT NULL"
    ).fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM ingest_run").fetchone()[0] == 1
    conn.close()


def test_hash_distinto_mismo_scope_supersede_e_inserta(planilla, tmp_path, db):
    """Caso 2 del contrato: contenido nuevo reemplaza la version vigente."""
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db)
    otro = _escribir(
        tmp_path / "jll_v2.xlsx",
        rr_rows=[_rr_row(total=99.0, tasa=0.99)],
    )
    r2 = commit(otro, db)
    assert r2["status"] == "ok"
    assert r2["superseded"] >= 2

    conn = get_conn_for(db)
    vivas = conn.execute(
        "SELECT renta_uf FROM raw_rent_roll_line WHERE superseded_at IS NULL"
    ).fetchall()
    assert [v[0] for v in vivas] == [99.0]
    conn.close()


def test_un_solo_snapshot_vivo_por_scope(planilla, tmp_path, db):
    """Tras dos ingestas del mismo grano no puede quedar mas de un snapshot."""
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db)
    commit(_escribir(tmp_path / "b.xlsx", [_rr_row(total=77.0)],
                     aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                                "Ingresos", 1, "X", "d", None, 5.0]],
                     cart_rows=[["Torre A", "C", "C", "1", "D", "RENTAS", "L",
                                 "2026-06-01", "2026-07-31", 1, 1.0, 0.0, 0.0,
                                 0.0, 0.0, 0.0, 1.0]],
                     rec_rows=[["Torre A", "2026-07-31", 9.0]]), db)

    conn = get_conn_for(db)
    for tabla in ("raw_rent_roll_line",) + repo_jll_v2.TABLAS:
        n = conn.execute(
            "SELECT COUNT(DISTINCT file_hash) FROM %s "
            " WHERE superseded_at IS NULL" % tabla
        ).fetchone()[0]
        assert n == 1, "%s quedo con %d snapshots vivos" % (tabla, n)
    conn.close()


def test_run_fallido_no_bloquea_reintentar_el_mismo_hash(planilla, db, monkeypatch):
    """Caso 3: fallo -> retry con el mismo hash -> exito."""
    import tools.db.ingest_jll_planilla as mod
    from tools.db.ingest_jll_planilla import commit

    real = mod.repo_jll_v2.insert_lines

    def explota(*a, **kw):
        raise RuntimeError("fallo simulado en la ultima hoja")

    monkeypatch.setattr(mod.repo_jll_v2, "insert_lines", explota)
    with pytest.raises(RuntimeError):
        commit(planilla, db)

    conn = get_conn_for(db)
    # Atomicidad: ni una fila, en ninguna tabla.
    for tabla in ("raw_rent_roll_line",) + repo_jll_v2.TABLAS:
        assert conn.execute("SELECT COUNT(*) FROM %s" % tabla).fetchone()[0] == 0
    # Pero el intento quedo registrado.
    estado = conn.execute(
        "SELECT status, error FROM ingest_run ORDER BY id DESC LIMIT 1"
    ).fetchone()
    assert estado[0] == "failed"
    assert "fallo simulado" in estado[1]
    conn.close()

    # Y el mismo archivo se puede reintentar.
    monkeypatch.setattr(mod.repo_jll_v2, "insert_lines", real)
    r = commit(planilla, db)
    assert r["status"] == "ok"

    conn = get_conn_for(db)
    assert conn.execute("SELECT COUNT(*) FROM raw_rent_roll_line").fetchone()[0] == 2
    estados = [r[0] for r in conn.execute(
        "SELECT status FROM ingest_run ORDER BY id")]
    assert estados == ["failed", "ok"]
    conn.close()


# ── semantica no mixta ──────────────────────────────────────────────────────

def test_no_se_puede_mezclar_semantica_de_renta(planilla, db):
    from tools.db.ingest_jll_planilla import commit, validate

    conn = get_conn_for(db)
    conn.execute(
        "INSERT INTO raw_rent_roll_line (activo_key, periodo, unidad, m2, renta_uf) "
        "VALUES ('Torre A', '2026-07', 'legacy', 50, 0.5)"
    )
    conn.commit()
    assert repo_rent_roll.semanticas_vivas(conn, "Torre A", "2026-07") == {
        "indeterminada"
    }
    conn.close()

    r = validate(planilla, db)
    assert r["puede_commitear"] is False
    assert r["conflictos_semantica"]

    with pytest.raises(ValueError, match="semantica de renta mezclada"):
        commit(planilla, db)


# ── lineage y procedencia ───────────────────────────────────────────────────

def test_filas_nuevas_declaran_procedencia(planilla, db):
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db)
    conn = get_conn_for(db)
    for tabla in ("raw_rent_roll_line",) + repo_jll_v2.TABLAS:
        faltan = conn.execute(
            "SELECT COUNT(*) FROM %s WHERE fuente_proveedor IS NULL "
            "   OR fuente_formato IS NULL" % tabla
        ).fetchone()[0]
        assert faltan == 0, "%s tiene filas sin procedencia" % tabla
        proveedor = [
            tuple(r) for r in conn.execute(
                "SELECT DISTINCT fuente_proveedor, fuente_formato FROM %s" % tabla
            )
        ]
        assert proveedor == [(FUENTE_PROVEEDOR, FUENTE_FORMATO)]
    conn.close()


def test_repo_rechaza_insert_sin_procedencia(db):
    conn = get_conn_for(db)
    with pytest.raises(repo_jll_v2.LineageIncompleto):
        repo_jll_v2.insert_lines(
            conn, "raw_recaudacion", [{"activo_key": "Torre A",
                                       "periodo": "2026-07", "monto_uf": 1.0}],
            ingest_run_id=1, fuente_proveedor="", fuente_formato="x",
            source_file="f", source_sheet="s", file_hash="h",
        )
    conn.close()


def test_rubro_sin_mapping_se_persiste_como_unmapped(planilla, db):
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db)
    conn = get_conn_for(db)
    filas = dict(conn.execute(
        "SELECT rubro, mapping_status FROM raw_movimiento_contable_line"
    ).fetchall())
    # raw-first: la fila existe aunque no tenga cuenta
    assert filas["Ingreso por intereses"] == "unmapped"
    assert filas["Ingreso por arriendo"] == "mapped"
    conn.close()


# -- derivacion a ER y lineage ----------------------------------------------

@pytest.fixture
def db_con_uf(db):
    conn = get_conn_for(db)
    conn.executemany(
        "INSERT OR IGNORE INTO raw_uf_diaria (fecha, valor, fuente) VALUES (?, ?, 'TEST')",
        [("2026-07-%02d" % d, 40000.0 + d) for d in range(1, 32)],
    )
    conn.commit()
    conn.close()
    return db


def test_er_derivado_tiene_lineage_a_los_movimientos(planilla, db_con_uf):
    from tools.db.ingest_jll_planilla import commit

    r = commit(planilla, db_con_uf)
    assert r["er_derivado"]["filas_jll_v2"] == 1  # solo el rubro mapeado

    conn = get_conn_for(db_con_uf)
    fila = conn.execute(
        "SELECT id, cuenta_codigo, monto_uf FROM raw_er_activo_line "
        " WHERE origen = 'jll_v2'"
    ).fetchone()
    assert fila["cuenta_codigo"] == "PT_ING_ARR"

    aportes = conn.execute(
        "SELECT aporte_uf FROM raw_er_movimiento_lineage WHERE er_line_id = ?",
        (fila["id"],),
    ).fetchall()
    assert aportes, "toda fila jll_v2 debe tener lineage"
    # el puente reproduce el agregado sin recomputarlo
    assert sum(a["aporte_uf"] for a in aportes) == pytest.approx(fila["monto_uf"])
    conn.close()


def test_reglas_internas_apuntan_a_su_version(planilla, db_con_uf):
    from tools.db.ingest_jll_planilla import commit

    commit(planilla, db_con_uf)
    conn = get_conn_for(db_con_uf)
    filas = conn.execute(
        "SELECT e.cuenta_codigo, r.version FROM raw_er_activo_line e "
        "  JOIN dim_er_regla_interna r ON r.id = e.origen_regla_id "
        " WHERE e.origen = 'regla_interna'"
    ).fetchall()
    assert {f["cuenta_codigo"] for f in filas} == {"PT_CONTRIB", "PT_SEG"}
    assert all(f["version"] == 1 for f in filas)
    conn.close()


def test_er_sin_uf_no_pierde_datos_crudos(planilla, db):
    """Sin UF no se puede derivar contribuciones, pero el raw entra igual."""
    from tools.db.ingest_jll_planilla import commit

    r = commit(planilla, db)
    assert r["status"] == "ok"
    assert r["er_derivado"]["reglas_no_evaluadas"]

    conn = get_conn_for(db)
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_movimiento_contable_line").fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM raw_rent_roll_line").fetchone()[0] == 2
    conn.close()


def test_manda_lo_interno_sobre_lo_de_jll(tmp_path, db_con_uf):
    """Contribuciones vienen en el archivo, pero el ER usa el valor interno."""
    from tools.db.ingest_jll_planilla import commit

    path = _escribir(
        tmp_path / "contrib.xlsx",
        rr_rows=[_rr_row()],
        aux_rows=[["Torre A", "2026-07-31", "Contribuciones", "Gastos", 8,
                   "SII", "Contribuciones", None, -999.0]],
    )
    r = commit(path, db_con_uf)

    alertas = r["er_derivado"]["alertas_control_cruzado"]
    assert len(alertas) == 1
    assert alertas[0]["cuenta_codigo"] == "PT_CONTRIB"
    assert alertas[0]["jll_uf"] == -999.0

    conn = get_conn_for(db_con_uf)
    filas = conn.execute(
        "SELECT origen, monto_uf FROM raw_er_activo_line "
        " WHERE cuenta_codigo = 'PT_CONTRIB'"
    ).fetchall()
    # una sola fila, y es la interna: el valor de JLL no se persistio en el ER
    assert len(filas) == 1
    assert filas[0]["origen"] == "regla_interna"
    assert filas[0]["monto_uf"] != -999.0
    conn.close()


# -- compatibilidad aguas abajo ---------------------------------------------

def test_vacancia_funciona_sobre_datos_v2(tmp_path, db):
    """Las vistas de vacancia deben leer filas v2 sin cambiar de forma.

    El modo de falla que esto previene es silencioso: sin la traduccion a
    extra_json.tipo_activo_2 y a arrendatario='Vacante', la vista clasificaria
    todo como 'Otro' y contaria cero vacantes.
    """
    from tools.db.ingest_jll_planilla import commit

    path = _escribir(tmp_path / "vac.xlsx", [
        _rr_row(local="1", categoria="OFICINA", estado="Ocupado", gla=100.0),
        _rr_row(local="2", categoria="OFICINA", estado="Vacante", gla=50.0,
                total=0, tasa=0, razon=None),
        _rr_row(local="E1", categoria="ESTACIONAMIENTO", estado="Vacante",
                gla=20.0, total=0, tasa=0, razon=None),
    ])
    commit(path, db)

    conn = get_conn_for(db)
    tipos = {
        r["tipo_unidad"]: (r["m2_gla"], r["m2_vacantes"])
        for r in conn.execute(
            "SELECT tipo_unidad, m2_gla, m2_vacantes FROM v_vacancia_activo_tipo "
            " WHERE activo_key='Torre A' AND periodo='2026-07'"
        )
    }
    assert tipos["Oficinas"] == (150.0, 50.0)
    assert tipos["Estacionamiento"] == (20.0, 20.0)

    total = conn.execute(
        "SELECT m2_gla, m2_vacantes, vacancia_pct FROM v_vacancia_activo "
        " WHERE activo_key='Torre A' AND periodo='2026-07'"
    ).fetchone()
    # el estacionamiento queda fuera del universo rentable
    assert total["m2_gla"] == 150.0
    assert total["m2_vacantes"] == 50.0
    assert total["vacancia_pct"] == pytest.approx(1 / 3)
    conn.close()


def test_ug_es_visible_y_no_se_diluye_en_otro(tmp_path, db):
    """UG queda como categoria propia mientras negocio no confirme su tratamiento."""
    from tools.db.ingest_jll_planilla import commit

    path = _escribir(tmp_path / "ug2.xlsx", [
        _rr_row(local="1", categoria="OFICINA", gla=100.0),
        _rr_row(local="U1", categoria="UG", gla=30.0),
    ])
    commit(path, db)

    conn = get_conn_for(db)
    tipos = {
        r["tipo_unidad"] for r in conn.execute(
            "SELECT tipo_unidad FROM v_vacancia_activo_tipo "
            " WHERE activo_key='Torre A' AND periodo='2026-07'"
        )
    }
    assert "UG" in tipos, "UG no debe diluirse en 'Otro'"
    assert "Otro" not in tipos
    conn.close()


def test_rent_roll_semantic_expone_filas_v2(tmp_path, db):
    from tools.db.ingest_jll_planilla import commit

    path = _escribir(tmp_path / "sem.xlsx", [_rr_row(categoria="LOCAL COMERCIAL")])
    commit(path, db)

    conn = get_conn_for(db)
    fila = conn.execute(
        "SELECT * FROM v_rent_roll_semantic WHERE activo_key='Torre A'"
    ).fetchone()
    assert fila is not None
    assert fila["unit_category"] == "local"
    assert fila["occupancy_status"] == "occupied"
    conn.close()


# -- regresiones del patch review -------------------------------------------

def test_scope_es_por_hoja_no_global(tmp_path, db):
    """Cada tabla supersede SOLO lo que su propia hoja cubre.

    Bug que esto previene: con un scope global (union de las cuatro hojas), un
    archivo B que trae auxiliar de 2026-01 y rent roll de 2026-07 superseder??a
    tambien el rent roll de 2026-01 cargado por A -- sin reemplazarlo, porque B
    no trae rent roll de ese mes. Borrado silencioso de datos vigentes.
    """
    from tools.db.ingest_jll_planilla import commit

    a = _escribir(
        tmp_path / "a.xlsx",
        rr_rows=[_rr_row(corte="2026-01-01", local="101")],
        aux_rows=[["Torre A", "2026-01-31", "Ingreso por arriendo", "Ingresos", 1,
                   "ACME", "d", None, 10.0]],
    )
    commit(a, db)

    b = _escribir(
        tmp_path / "b.xlsx",
        rr_rows=[_rr_row(corte="2026-07-01", local="102")],
        aux_rows=[["Torre A", "2026-01-31", "Ingreso por arriendo", "Ingresos", 1,
                   "ACME", "d", None, 20.0]],
    )
    commit(b, db)

    conn = get_conn_for(db)
    vivos = dict(conn.execute(
        "SELECT periodo, COUNT(*) FROM raw_rent_roll_line "
        " WHERE superseded_at IS NULL GROUP BY periodo"
    ).fetchall())
    # el rent roll de 2026-01 que trajo A sigue vivo: B no lo reemplaza
    assert vivos.get("2026-01") == 1, "B superseder??a rent roll que no reemplaza"
    assert vivos.get("2026-07") == 1
    # y el auxiliar de 2026-01 SI fue reemplazado por B
    montos = [r[0] for r in conn.execute(
        "SELECT monto_uf FROM raw_movimiento_contable_line "
        " WHERE periodo='2026-01' AND superseded_at IS NULL")]
    assert montos == [20.0]
    conn.close()


def test_er_derivado_deja_una_sola_fila_viva_por_cuenta(tmp_path, db_con_uf):
    """A -> B debe dejar exactamente una fila de ER viva y lineage solo hacia B.

    Sin supersession del ER, cada ingesta agregaba una fila mas para la misma
    cuenta y el NOI quedaba duplicado.
    """
    from tools.db.ingest_jll_planilla import commit

    a = _escribir(tmp_path / "era.xlsx", rr_rows=[_rr_row()],
                  aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                             "Ingresos", 1, "ACME", "d", None, 60.0]])
    commit(a, db_con_uf)
    b = _escribir(tmp_path / "erb.xlsx", rr_rows=[_rr_row()],
                  aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                             "Ingresos", 1, "ACME", "d", None, 99.0]])
    commit(b, db_con_uf)

    conn = get_conn_for(db_con_uf)
    vivas = conn.execute(
        "SELECT id, monto_uf FROM raw_er_activo_line "
        " WHERE activo_key='Torre A' AND periodo='2026-07' "
        "   AND cuenta_codigo='PT_ING_ARR' AND superseded_at IS NULL"
    ).fetchall()
    assert len(vivas) == 1, "quedaron %d filas vivas: NOI duplicado" % len(vivas)
    assert vivas[0]["monto_uf"] == 99.0

    # el lineage de la fila viva apunta solo a movimientos vivos (los de B)
    origenes = conn.execute(
        """SELECT m.monto_uf, m.superseded_at FROM raw_er_movimiento_lineage l
             JOIN raw_movimiento_contable_line m ON m.id = l.movimiento_id
            WHERE l.er_line_id = ?""",
        (vivas[0]["id"],),
    ).fetchall()
    assert [r["monto_uf"] for r in origenes] == [99.0]
    assert all(r["superseded_at"] is None for r in origenes)
    conn.close()


def test_er_no_supersede_cuentas_que_el_archivo_no_reconstruye(tmp_path, db_con_uf):
    """El reemplazo es quirurgico: una cuenta ajena al archivo sigue viva."""
    from tools.db.ingest_jll_planilla import commit

    conn = get_conn_for(db_con_uf)
    conn.execute(
        "INSERT INTO raw_er_activo_line (activo_key, periodo, cuenta_codigo, monto_clp) "
        "VALUES ('Torre A', '2026-07', 'PT_GAST_ADIC', -500)"
    )
    conn.commit()
    conn.close()

    path = _escribir(tmp_path / "q.xlsx", rr_rows=[_rr_row()],
                     aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                                "Ingresos", 1, "ACME", "d", None, 60.0]])
    commit(path, db_con_uf)

    conn = get_conn_for(db_con_uf)
    ajena = conn.execute(
        "SELECT COUNT(*) FROM raw_er_activo_line WHERE cuenta_codigo='PT_GAST_ADIC' "
        "  AND superseded_at IS NULL"
    ).fetchone()[0]
    assert ajena == 1, "se supersedio una cuenta que el archivo no puede reponer"
    conn.close()


def test_reglas_internas_no_hacen_producto_cartesiano(tmp_path, db_con_uf):
    """Solo se evaluan los pares (activo, periodo) que el auxiliar cubre."""
    from tools.db.ingest_jll_planilla import commit

    path = _escribir(tmp_path / "cart.xlsx", rr_rows=[_rr_row()],
                     aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                                "Ingresos", 1, "ACME", "d", None, 60.0]])
    r = commit(path, db_con_uf)

    conn = get_conn_for(db_con_uf)
    periodos = {p[0] for p in conn.execute(
        "SELECT DISTINCT periodo FROM raw_er_activo_line WHERE origen='regla_interna'")}
    assert periodos == {"2026-07"}, "se generaron periodos que el auxiliar no cubre"
    # 2 reglas vigentes para Torre A (PT_CONTRIB + PT_SEG), un solo periodo
    assert r["er_derivado"]["filas_regla_interna"] == 2
    conn.close()


def test_hash_exitoso_sigue_siendo_duplicado_tras_ser_supersedido(tmp_path, db):
    """A(ok) -> B(ok, supersede A) -> A otra vez debe ser NO-OP.

    Con el chequeo por filas vivas, A dejaba de contar como duplicado apenas B
    lo superseder??a: reingestarlo habria superseder??o a B y hecho retroceder
    los datos.
    """
    from tools.db.ingest_jll_planilla import commit

    a = _escribir(tmp_path / "ha.xlsx", rr_rows=[_rr_row(total=11.0)])
    b = _escribir(tmp_path / "hb.xlsx", rr_rows=[_rr_row(total=22.0)])
    assert commit(a, db)["status"] == "ok"
    assert commit(b, db)["status"] == "ok"

    conn = get_conn_for(db)
    assert [r[0] for r in conn.execute(
        "SELECT renta_uf FROM raw_rent_roll_line WHERE superseded_at IS NULL")] == [22.0]
    conn.close()

    # A ya se ingesto con exito alguna vez: no debe volver a entrar
    assert commit(a, db)["status"] == "skipped_duplicate"

    conn = get_conn_for(db)
    vivas = [r[0] for r in conn.execute(
        "SELECT renta_uf FROM raw_rent_roll_line WHERE superseded_at IS NULL")]
    assert vivas == [22.0], "B debe seguir vivo; A no puede resucitar"
    conn.close()


def test_fila_sin_monto_bloquea_el_commit(tmp_path, db):
    """fail-loud: sin monto la fila no es persistible, no se descarta callada."""
    from tools.db.ingest_jll_planilla import commit, validate

    path = _escribir(tmp_path / "nm.xlsx", rr_rows=[_rr_row()],
                     aux_rows=[["Torre A", "2026-07-31", "Ingreso por arriendo",
                                "Ingresos", 1, "ACME", "d", None, None]])
    r = validate(path, db)
    assert r["puede_commitear"] is False
    assert "sin_monto" in r["anomalias"]["por_tipo"]
    assert commit(path, db)["status"] == "bloqueado"

    conn = get_conn_for(db)
    assert conn.execute("SELECT COUNT(*) FROM raw_rent_roll_line").fetchone()[0] == 0
    conn.close()


def test_fila_sin_rubro_bloquea_el_commit(tmp_path, db):
    from tools.db.ingest_jll_planilla import validate

    path = _escribir(tmp_path / "nr.xlsx", rr_rows=[_rr_row()],
                     aux_rows=[["Torre A", "2026-07-31", None, "Ingresos", 1,
                                "ACME", "d", None, 5.0]])
    r = validate(path, db)
    assert r["puede_commitear"] is False
    assert "sin_rubro" in r["anomalias"]["por_tipo"]


def test_historico_jll_v1_que_se_va_a_superseder_no_bloquea(tmp_path, db):
    """El gate evalua el estado POST-supersession.

    Tras la migracion 089 el historico JLL v1 tiene fuente_proveedor='JLL', asi
    que el propio commit lo supersede. Bloquear por el seria impedir el cutover
    que este pipeline existe para hacer.
    """
    from tools.db.ingest_jll_planilla import commit, validate

    conn = get_conn_for(db)
    conn.execute(
        "INSERT INTO raw_rent_roll_line (activo_key, periodo, unidad, m2, renta_uf, "
        "  renta_semantica, fuente_proveedor, fuente_formato) "
        "VALUES ('Torre A', '2026-07', 'v1', 50, 0.5, 'indeterminada', 'JLL', 'jll_v1')"
    )
    conn.commit()
    conn.close()

    path = _escribir(tmp_path / "cut.xlsx", rr_rows=[_rr_row()])
    r = validate(path, db)
    assert r["conflictos_semantica"] == [], "el historico JLL v1 no debe bloquear"
    assert r["puede_commitear"] is True
    assert commit(path, db)["status"] == "ok"

    conn = get_conn_for(db)
    sem = {s[0] for s in conn.execute(
        "SELECT DISTINCT renta_semantica FROM raw_rent_roll_line "
        " WHERE periodo='2026-07' AND superseded_at IS NULL")}
    assert sem == {"total_uf"}, "no puede quedar semantica mezclada"
    conn.close()


def test_otra_procedencia_si_bloquea(tmp_path, db):
    """Lo que el barrido NO alcanza si debe bloquear: procedencia distinta."""
    from tools.db.ingest_jll_planilla import validate

    conn = get_conn_for(db)
    conn.execute(
        "INSERT INTO raw_rent_roll_line (activo_key, periodo, unidad, m2, renta_uf, "
        "  renta_semantica, fuente_proveedor) "
        "VALUES ('Torre A', '2026-07', 'x', 50, 0.5, 'indeterminada', 'TresA')"
    )
    conn.commit()
    conn.close()

    path = _escribir(tmp_path / "res.xlsx", rr_rows=[_rr_row()])
    r = validate(path, db)
    assert r["conflictos_semantica"], "semantica residual debe bloquear"
    assert r["puede_commitear"] is False


def test_source_file_conserva_el_nombre_real_no_el_temporal(tmp_path, db):
    """La DB y el ingest_run guardan el nombre del upload, no `jll_v2_xxx.xlsx`."""
    from tools.db.ingest_jll_planilla import commit

    ruta_lectura = _escribir(tmp_path / "jll_v2_ab12cd.xlsx", rr_rows=[_rr_row()])
    commit(ruta_lectura, db, source_name="2607 Datos Carga Rentas Toesca.xlsx")

    conn = get_conn_for(db)
    nombres = {r[0] for r in conn.execute(
        "SELECT DISTINCT source_file FROM raw_rent_roll_line WHERE superseded_at IS NULL")}
    assert nombres == {"2607 Datos Carga Rentas Toesca.xlsx"}
    run = conn.execute(
        "SELECT source_file FROM ingest_run ORDER BY id DESC LIMIT 1").fetchone()[0]
    assert run == "2607 Datos Carga Rentas Toesca.xlsx"
    conn.close()


def test_repo_v2_falla_ante_violacion_de_restriccion(db):
    """El pipeline v2 usa INSERT normal: una restriccion incumplida ABORTA.

    Con `INSERT OR IGNORE` la segunda fila se descartaba en silencio y la
    ingesta reportaba exito habiendo perdido un dato. Este test ejercita el
    INSERT directamente, porque las anomalias de `validate()` nunca dejan llegar
    una fila mal formada hasta la DB.
    """
    import sqlite3

    conn = get_conn_for(db)
    fila = {"activo_key": "Torre A", "periodo": "2026-07", "monto_uf": 1.0,
            "source_row": 2}
    comun = dict(ingest_run_id=None, fuente_proveedor="JLL", fuente_formato="jll_v2",
                 source_file="f.xlsx", source_sheet="Recaudado", file_hash="hash-x")

    repo_jll_v2.insert_lines(conn, "raw_recaudacion", [fila], **comun)
    # misma (file_hash, source_row): viola uq_recaudacion_vivo
    with pytest.raises(sqlite3.IntegrityError):
        repo_jll_v2.insert_lines(conn, "raw_recaudacion", [dict(fila, monto_uf=2.0)],
                                 **comun)
    conn.rollback()
    conn.close()


def test_repo_rent_roll_v2_falla_ante_violacion_de_restriccion(db):
    import sqlite3

    conn = get_conn_for(db)
    fila = {"activo_key": "Torre A", "periodo": "2026-07", "unidad": "1",
            "m2": 10.0, "renta_uf": 1.0, "renta_semantica": "total_uf",
            "source_row": 2, "file_hash": "hash-y", "source_file": "f.xlsx"}
    repo_rent_roll.insert_lines_v2(conn, [fila], ingest_run_id=None,
                                   fuente_proveedor="JLL", fuente_formato="jll_v2")
    with pytest.raises(sqlite3.IntegrityError):
        repo_rent_roll.insert_lines_v2(conn, [dict(fila, renta_uf=2.0)],
                                       ingest_run_id=None,
                                       fuente_proveedor="JLL", fuente_formato="jll_v2")
    conn.rollback()
    conn.close()


# -- patch review final: 4 fixes required -----------------------------------

def test_fail_closed_no_cae_a_jll_si_regla_interna_no_evaluable(tmp_path, db):
    """Sin UF, PT_CONTRIB no se puede evaluar internamente. JLL NO debe tomar
    el control por default: la cuenta queda solo en raw, el ER anterior no se
    toca, y el bloqueo se reporta explicitamente.
    """
    from tools.db.ingest_jll_planilla import commit

    conn = get_conn_for(db)
    conn.execute(
        "INSERT INTO raw_er_activo_line (activo_key, periodo, cuenta_codigo, monto_clp) "
        "VALUES ('Torre A', '2026-07', 'PT_CONTRIB', -999)"
    )
    conn.commit()
    conn.close()

    # `db` (sin db_con_uf) no tiene UF cargada: la formula de contribuciones
    # no se puede evaluar para 2026-07.
    path = _escribir(tmp_path / "fc.xlsx", rr_rows=[_rr_row()],
                     aux_rows=[["Torre A", "2026-07-31", "Contribuciones",
                                "Gastos", 8, "SII", "d", None, -500.0]])
    r = commit(path, db)
    assert r["status"] == "ok"

    bloqueadas = r["er_derivado"]["autoridad_interna_no_evaluable"]
    assert len(bloqueadas) == 1
    assert bloqueadas[0]["cuenta_codigo"] == "PT_CONTRIB"
    assert bloqueadas[0]["jll_uf_en_raw"] == -500.0

    conn = get_conn_for(db)
    # el ER anterior sigue vivo, tal cual, no se superseded ni se reemplazo
    vivas = conn.execute(
        "SELECT monto_clp, superseded_at FROM raw_er_activo_line "
        " WHERE activo_key='Torre A' AND periodo='2026-07' AND cuenta_codigo='PT_CONTRIB'"
    ).fetchall()
    assert len(vivas) == 1
    assert vivas[0]["monto_clp"] == -999
    assert vivas[0]["superseded_at"] is None
    # el valor de JLL SI quedo persistido en crudo, mapeado
    mov = conn.execute(
        "SELECT monto_uf, mapping_status FROM raw_movimiento_contable_line "
        " WHERE rubro='Contribuciones' AND activo_key='Torre A'"
    ).fetchone()
    assert mov["monto_uf"] == -500.0
    assert mov["mapping_status"] == "mapped"
    conn.close()


def test_apo3001_seguro_no_esta_en_cuentas_internas(tmp_path, db_con_uf):
    """La autoridad de APO3001_SEG es JLL, no interna: no debe fail-closed."""
    from tools.db import derive_er_jll_v2 as der

    assert "APO3001_SEG" not in der.CUENTAS_INTERNAS

    conn = get_conn_for(db_con_uf)
    conn.execute(
        "INSERT OR IGNORE INTO dim_activo (activo_key, nombre, fondo_key) "
        "VALUES ('Apo3001', 'Apoquindo 3001', 'TRI')"
    )
    from tools.db import repo_audit
    run_id = repo_audit.start_ingest_run(conn, "test", "f.xlsx", "hseg")
    conn.execute(
        "INSERT INTO raw_movimiento_contable_line "
        "  (activo_key, periodo, rubro, clasificacion, monto_uf, "
        "   cuenta_codigo_mapeada, mapping_status, fuente_proveedor, "
        "   fuente_formato, source_file, source_sheet, source_row, file_hash, "
        "   ingest_run_id) "
        "VALUES ('Apo3001', '2026-07', 'Seguros', 'Gastos', -5.0, "
        "  'APO3001_SEG', 'mapped', 'JLL', 'jll_v2', 'f.xlsx', 'AuxiliarContable', 1, "
        "  'hseg', ?)",
        (run_id,),
    )
    conn.commit()

    er = der.derivar(conn, file_hash="hseg", source_file="f.xlsx",
                     ingest_run_id=run_id, periodos=["2026-07"])
    assert er["autoridad_interna_no_evaluable"] == []
    assert er["filas_jll_v2"] >= 1

    fila = conn.execute(
        "SELECT monto_uf, origen FROM raw_er_activo_line "
        " WHERE cuenta_codigo='APO3001_SEG' AND superseded_at IS NULL"
    ).fetchone()
    assert fila["monto_uf"] == -5.0
    assert fila["origen"] == "jll_v2"
    conn.close()


def test_pares_incluyen_activo_periodo_solo_con_rubros_unmapped(tmp_path, db_con_uf):
    """Apo3001 con TODOS sus rubros unmapped igual evalua sus reglas internas.

    'Ingreso por arriendo' de Apo3001 queda unmapped a proposito (dos cuentas
    de ingreso posibles, sin criterio de desambiguacion). Un periodo cuyo unico
    movimiento sea ese no debe quedar sin ER de contribuciones.
    """
    from tools.db.ingest_jll_planilla import commit

    conn = get_conn_for(db_con_uf)
    conn.execute(
        "INSERT OR IGNORE INTO dim_activo (activo_key, nombre, fondo_key) "
        "VALUES ('Apo3001', 'Apoquindo 3001', 'TRI')"
    )
    conn.commit()
    conn.close()

    path = _escribir(
        tmp_path / "unm.xlsx",
        rr_rows=[_rr_row(portafolio="Apoquindo 3001")],
        aux_rows=[["Apoquindo 3001", "2026-07-31", "Ingreso por arriendo",
                   "Ingresos", 1, "X", "d", None, 100.0]],
    )
    r = commit(path, db_con_uf)
    assert r["status"] == "ok"

    conn = get_conn_for(db_con_uf)
    # el movimiento quedo unmapped, tal como se espera
    mapping = conn.execute(
        "SELECT mapping_status FROM raw_movimiento_contable_line "
        " WHERE activo_key='Apo3001' AND periodo='2026-07'"
    ).fetchone()
    assert mapping["mapping_status"] == "unmapped"

    # y AUN ASI la regla interna de contribuciones para Apo3001 se evaluo
    interna = conn.execute(
        "SELECT cuenta_codigo FROM raw_er_activo_line "
        " WHERE activo_key='Apo3001' AND periodo='2026-07' AND origen='regla_interna'"
    ).fetchall()
    assert {r["cuenta_codigo"] for r in interna} == {"APO3001_CONTRIB_SOBRETASA"}
    conn.close()


def test_aceptar_anomalias_no_puede_saltarse_conflicto_de_semantica(planilla, db):
    """Invariante de integridad, no anomalia de calidad: aceptar_anomalias no
    lo debe poder overridear jamas."""
    from tools.db.ingest_jll_planilla import commit

    conn = get_conn_for(db)
    conn.execute(
        "INSERT INTO raw_rent_roll_line (activo_key, periodo, unidad, m2, renta_uf) "
        "VALUES ('Torre A', '2026-07', 'legacy', 50, 0.5)"
    )
    conn.commit()
    conn.close()

    with pytest.raises(ValueError, match="semantica"):
        commit(planilla, db, aceptar_anomalias=True)

    conn = get_conn_for(db)
    # nada se escribio: la excepcion se lanzo ANTES del BEGIN IMMEDIATE
    assert conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE unidad != 'legacy'"
    ).fetchone()[0] == 0
    conn.close()


def test_idempotencia_sobrevive_a_crash_antes_de_finish_ingest_run(planilla, db):
    """Simula: los datos se commitearon (conn.commit() de la transaccion de
    datos corrio bien) pero el proceso murio ANTES de finish_ingest_run()
    (que corre en una conexion separada). El run queda 'started' para
    siempre. Un retry con el mismo archivo debe ser NO-OP, no debe duplicar.
    """
    from tools.db.ingest_jll_planilla import commit

    r1 = commit(planilla, db)
    assert r1["status"] == "ok"

    # Simular el crash: revertir el status del run a lo que habria quedado si
    # finish_ingest_run() nunca hubiera corrido.
    conn = get_conn_for(db)
    conn.execute(
        "UPDATE ingest_run SET status='started', ended_at=NULL WHERE id=?",
        (r1["ingest_run_id"],),
    )
    conn.commit()
    conn.close()

    conn = get_conn_for(db)
    antes = conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE superseded_at IS NULL"
    ).fetchone()[0]
    conn.close()

    r2 = commit(planilla, db)
    assert r2["status"] == "skipped_duplicate"

    conn = get_conn_for(db)
    despues = conn.execute(
        "SELECT COUNT(*) FROM raw_rent_roll_line WHERE superseded_at IS NULL"
    ).fetchone()[0]
    assert despues == antes, "el retry no debe duplicar ni modificar nada"
    conn.close()
