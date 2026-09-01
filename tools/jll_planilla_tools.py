"""Parser de la planilla unica de JLL ("JLL v2").

JLL dejo de entregar el `{AAMM} Rent Roll y NOI.xlsx` clasico. Ahora entrega un
archivo multi-hoja y multi-periodo con rent roll, auxiliar contable
(facturacion por tercero), cartera de morosos y recaudacion.

Este modulo solo LEE y NORMALIZA: no toca la DB. La persistencia vive en
tools/db/ingest_jll_planilla.py.

Principio rector: **fail-loud**. Ninguna fila se descarta en silencio. Todo lo
que no se pueda interpretar sale en `anomalias`, con su numero de fila de
origen, para que `validate()` lo muestre antes de cualquier commit.

JLL es *una fuente*, no la definicion de estos conceptos: los nombres de salida
son canonicos y otra fuente podria producirlos igual.
"""
from __future__ import annotations

import datetime as _dt
import os
import shutil
import tempfile
from dataclasses import dataclass, field

import openpyxl

FUENTE_PROVEEDOR = "JLL"
FUENTE_FORMATO = "jll_v2"

HOJA_RENT_ROLL = "RentRoll"
HOJA_AUXILIAR = "AuxiliarContable"
HOJA_CARTERA = "Cartera"
HOJA_RECAUDADO = "Recaudado"

# Hojas cuya presencia simultanea identifica el formato v2.
HOJAS_FIRMA = (HOJA_RENT_ROLL, HOJA_AUXILIAR)

# `portafolio` de JLL -> activo_key de dim_activo. Mapeo 1:1 y estable; todos
# existen en dim_activo. Un portafolio desconocido es una anomalia, no un
# descarte silencioso.
PORTAFOLIO_A_ACTIVO = {
    "Apoquindo 4700": "Apo4700",
    "Apoquindo 4501": "Apo4501",
    "Apoquindo 3001": "Apo3001",
    "Boulevard PT": "Boulevard",
    "Torre A": "Torre A",
}

ESTADOS_RENT_ROLL_VALIDOS = {"Vacante", "Ocupado"}


@dataclass
class Anomalia:
    hoja: str
    fila: int
    tipo: str
    detalle: str
    valor: object = None


@dataclass
class HojaParseada:
    hoja: str
    filas: list[dict] = field(default_factory=list)
    anomalias: list[Anomalia] = field(default_factory=list)
    filas_leidas: int = 0


@dataclass
class PlanillaParseada:
    source_file: str
    rent_roll: HojaParseada
    movimientos: HojaParseada
    cartera: HojaParseada
    recaudacion: HojaParseada

    @property
    def hojas(self) -> list[HojaParseada]:
        return [self.rent_roll, self.movimientos, self.cartera, self.recaudacion]

    @property
    def anomalias(self) -> list[Anomalia]:
        return [a for h in self.hojas for a in h.anomalias]

    def scope_de(self, hoja: HojaParseada) -> set[tuple[str, str, str]]:
        """Combinaciones (fuente_proveedor, activo_key, periodo) de UNA hoja.

        Cada tabla destino supersede sólo lo que su propia hoja cubre. Usar un
        scope global sería incorrecto: las hojas tienen coberturas temporales
        distintas (en el archivo del 2026-08-12 el auxiliar arranca en 2025-01 y
        el rent roll en 2025-11), así que un scope unificado superseder??a rent
        roll de períodos que el archivo no trae en esa hoja, borrando datos
        vigentes sin reemplazo.
        """
        return {
            (FUENTE_PROVEEDOR, fila["activo_key"], fila["periodo"])
            for fila in hoja.filas
            if fila["activo_key"] and fila["periodo"]
        }

    def scope_por_hoja(self) -> dict[str, set[tuple[str, str, str]]]:
        """{nombre_hoja: scope}. Lo que consume commit() para superseder."""
        return {hoja.hoja: self.scope_de(hoja) for hoja in self.hojas}

    def scope(self) -> set[tuple[str, str, str]]:
        """Unión de los scopes de todas las hojas.

        Sólo para reportar en `validate()` qué cubre el archivo en total. **No
        usar para superseder** — ver `scope_de`.
        """
        out: set[tuple[str, str, str]] = set()
        for hoja in self.hojas:
            out |= self.scope_de(hoja)
        return out

    def periodos(self) -> list[str]:
        return sorted({p for _, _, p in self.scope()})


# -- lectura de bajo nivel ---------------------------------------------------

def _abrir(filepath: str):
    """Abre el libro; si esta bloqueado por Excel, trabaja sobre una copia."""
    try:
        return openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), os.path.basename(filepath))
        shutil.copy2(filepath, tmp)
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True)


def es_formato_v2(filepath: str) -> bool:
    """True si el archivo tiene la firma de hojas del formato JLL v2."""
    wb = _abrir(filepath)
    try:
        return all(h in wb.sheetnames for h in HOJAS_FIRMA)
    finally:
        wb.close()


def _col_map(header: tuple) -> dict:
    """{nombre_normalizado: indice}. Mapea por nombre, nunca por posicion."""
    out = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        name = str(h).strip()
        if name and name not in out:
            out[name] = i
    return out


def _get(row: tuple, cols: dict, nombre: str):
    i = cols.get(nombre)
    if i is None or i >= len(row):
        return None
    return row[i]


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    try:
        return float(s.replace(".", "").replace(",", ".")) if "," in s else float(s)
    except ValueError:
        return None


def _texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _fecha_fuente(v):
    """Fecha original como texto ISO, sin truncar."""
    if v is None:
        return None
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return s or None


def _periodo(v):
    """Periodo canonico YYYY-MM.

    Convencion acordada 2026-08-28: el dia 1 identifica el mes en curso
    (2026-07-01 -> 2026-07). La serie del archivo (nov y dic con etiqueta de fin
    de mes, luego ene..jul con etiqueta de dia 1) da 9 meses consecutivos sin
    repetir diciembre solo bajo esta lectura. PENDIENTE de confirmacion con JLL;
    por eso cada fila conserva ademas su `fecha_corte_fuente`.
    """
    iso = _fecha_fuente(v)
    return iso[:7] if iso and len(iso) >= 7 else None


def _resolver_activo(portafolio, hoja: str, nfila: int, anomalias: list):
    nombre = _texto(portafolio)
    if not nombre:
        anomalias.append(Anomalia(hoja, nfila, "portafolio_vacio", "Fila sin portafolio"))
        return None
    activo = PORTAFOLIO_A_ACTIVO.get(nombre)
    if activo is None:
        anomalias.append(
            Anomalia(hoja, nfila, "portafolio_desconocido",
                     "'%s' no esta en PORTAFOLIO_A_ACTIVO" % nombre, nombre)
        )
    return activo


# -- hojas -------------------------------------------------------------------

def parse_rent_roll(rows: list) -> HojaParseada:
    hoja = HojaParseada(HOJA_RENT_ROLL)
    if not rows:
        return hoja
    cols = _col_map(rows[0])

    for offset, row in enumerate(rows[1:], start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        hoja.filas_leidas += 1
        activo = _resolver_activo(
            _get(row, cols, "portafolio"), HOJA_RENT_ROLL, offset, hoja.anomalias
        )

        corte = _get(row, cols, "fecha_corte_rent_roll")
        periodo = _periodo(corte)
        if periodo is None:
            hoja.anomalias.append(
                Anomalia(HOJA_RENT_ROLL, offset, "sin_fecha_corte",
                         "fecha_corte_rent_roll vacia o ilegible", corte)
            )

        estado = _texto(_get(row, cols, "estado_rent_roll"))
        if estado is not None and estado not in ESTADOS_RENT_ROLL_VALIDOS:
            # Sintoma conocido de columnas desalineadas en el origen: aparecen
            # codigos de local ('1-TA', 'BS4B07', '2009-E') donde deberia ir el
            # estado. ~1.040 filas en el archivo del 2026-08-12.
            hoja.anomalias.append(
                Anomalia(HOJA_RENT_ROLL, offset, "estado_invalido",
                         "estado_rent_roll='%s' no es Vacante/Ocupado" % estado, estado)
            )

        categoria = _texto(_get(row, cols, "categoria_general")) or _texto(
            _get(row, cols, "categoría_general")
        )
        if categoria is None:
            hoja.anomalias.append(
                Anomalia(HOJA_RENT_ROLL, offset, "sin_categoria", "categoria_general vacia")
            )
        elif categoria.upper() == "UG":
            # Tratamiento pendiente de confirmacion con negocio: NO se asume que
            # deba excluirse junto con los estacionamientos.
            hoja.anomalias.append(
                Anomalia(HOJA_RENT_ROLL, offset, "categoria_ug_pendiente",
                         "categoria 'UG' sin tratamiento confirmado", categoria)
            )

        m2 = _num(_get(row, cols, "informacion_locatario_gla")) or _num(
            _get(row, cols, "información_locatario_gla")
        )
        total = _num(_get(row, cols, "renta_renta_real"))
        tasa = _num(_get(row, cols, "renta_m2"))

        # En v2 `renta_renta_real` se computa como tasa x area (99,7% de las
        # filas evaluables), o sea es el total CONTRACTUAL pese al nombre. Ver
        # docs/rent-roll-renta-semantics-v1.md. Cuando ambos vienen y no
        # reconcilian, es un dato malo del origen, no una interpretacion nuestra.
        if m2 and tasa and total and abs(tasa * m2 - total) / max(abs(total), 1e-9) > 0.01:
            hoja.anomalias.append(
                Anomalia(HOJA_RENT_ROLL, offset, "renta_no_reconcilia",
                         "tasa*m2=%.2f vs total=%.2f" % (tasa * m2, total),
                         {"tasa": tasa, "m2": m2, "total": total})
            )

        hoja.filas.append({
            "activo_key": activo,
            "periodo": periodo,
            "fecha_corte_fuente": _fecha_fuente(corte),
            "unidad": _texto(_get(row, cols, "local")),
            "arrendatario": (
                _texto(_get(row, cols, "informacion_locatario_razon_social"))
                or _texto(_get(row, cols, "información_locatario_razón_social"))
            ),
            "m2": m2,
            "renta_uf": total,          # total UF (target de la migracion 085)
            "renta_uf_m2": tasa,        # tasa UF/m2
            "renta_semantica": "total_uf",
            "vencimiento": _fecha_fuente(
                _get(row, cols, "fechas_de_ocupacion_fin")
                or _get(row, cols, "fechas_de_ocupación_fin")
            ),
            "estado_rent_roll": estado,
            "categoria_general": categoria,
            "piso": _texto(_get(row, cols, "piso")),
            "marca": (
                _texto(_get(row, cols, "informacion_locatario_marca"))
                or _texto(_get(row, cols, "información_locatario_marca"))
            ),
            "contrato": _texto(_get(row, cols, "#Contrato")),
            "estado_contrato": (
                _texto(_get(row, cols, "informacion_locatario_estado_contrato"))
                or _texto(_get(row, cols, "información_locatario_estado_contrato"))
            ),
            "fecha_inicio": _fecha_fuente(
                _get(row, cols, "fechas_de_ocupacion_inicio")
                or _get(row, cols, "fechas_de_ocupación_inicio")
            ),
            "source_row": offset,
        })
    return hoja


def parse_auxiliar_contable(rows: list) -> HojaParseada:
    hoja = HojaParseada(HOJA_AUXILIAR)
    if not rows:
        return hoja
    cols = _col_map(rows[0])

    for offset, row in enumerate(rows[1:], start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        hoja.filas_leidas += 1
        activo = _resolver_activo(
            _get(row, cols, "portafolio"), HOJA_AUXILIAR, offset, hoja.anomalias
        )
        mes = _get(row, cols, "mes")
        periodo = _periodo(mes)
        if periodo is None:
            hoja.anomalias.append(
                Anomalia(HOJA_AUXILIAR, offset, "sin_periodo", "mes vacio o ilegible", mes)
            )

        # Solo `credito` viene poblada en el archivo del 2026-08-12; debito y
        # los saldos estan 100% vacios. Se lee credito y se deja constancia si
        # apareciera un debito, que cambiaria la interpretacion del signo.
        monto = _num(_get(row, cols, "credito")) or _num(_get(row, cols, "crédito"))
        debito = _num(_get(row, cols, "debito")) or _num(_get(row, cols, "débito"))
        if debito:
            hoja.anomalias.append(
                Anomalia(HOJA_AUXILIAR, offset, "debito_inesperado",
                         "la columna debito trae valor; revisar convencion de signo", debito)
            )
        if monto is None:
            hoja.anomalias.append(
                Anomalia(HOJA_AUXILIAR, offset, "sin_monto", "credito vacio")
            )

        rubro = _texto(_get(row, cols, "rubro_presupuestal"))
        if rubro is None:
            hoja.anomalias.append(
                Anomalia(HOJA_AUXILIAR, offset, "sin_rubro", "rubro_presupuestal vacio")
            )

        hoja.filas.append({
            "activo_key": activo,
            "periodo": periodo,
            "fecha_fuente": _fecha_fuente(mes),
            "rubro": rubro,
            "clasificacion": (
                _texto(_get(row, cols, "clasificacion presupuestal"))
                or _texto(_get(row, cols, "clasificación presupuestal"))
            ),
            "codigo_rubro": (
                _texto(_get(row, cols, "codigo_contable"))
                or _texto(_get(row, cols, "código_contable"))
            ),
            "tercero": _texto(_get(row, cols, "nombre_del_tercero")),
            "descripcion": (
                _texto(_get(row, cols, "descripcion"))
                or _texto(_get(row, cols, "descripción"))
            ),
            "monto_uf": monto,
            "source_row": offset,
        })
    return hoja


def parse_cartera(rows: list) -> HojaParseada:
    hoja = HojaParseada(HOJA_CARTERA)
    if not rows:
        return hoja
    cols = _col_map(rows[0])

    for offset, row in enumerate(rows[1:], start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        hoja.filas_leidas += 1
        # La cabecera trae 'Portafolio ' con espacio final; _col_map ya hace strip.
        activo = _resolver_activo(
            _get(row, cols, "Portafolio"), HOJA_CARTERA, offset, hoja.anomalias
        )
        corte = _get(row, cols, "fecha_de_corte")
        periodo = _periodo(corte)
        if periodo is None:
            hoja.anomalias.append(
                Anomalia(HOJA_CARTERA, offset, "sin_fecha_corte",
                         "fecha_de_corte vacia o ilegible", corte)
            )

        buckets = {
            "vencido_1_30": _num(_get(row, cols, "vencido_1_a_30")),
            "vencido_31_60": _num(_get(row, cols, "vencido_31_a_60")),
            "vencido_61_90": _num(_get(row, cols, "vencido_61_a_90")),
            "vencido_mas_91": (
                _num(_get(row, cols, "vencido_mas_de_91"))
                or _num(_get(row, cols, "vencido_más_de_91"))
            ),
            "saldo_por_vencer": _num(_get(row, cols, "saldo_por_vencer")),
        }
        total = _num(_get(row, cols, "total_cartera"))
        suma = sum(v for v in buckets.values() if v is not None)
        if total is not None and abs(suma - total) > 0.01:
            hoja.anomalias.append(
                Anomalia(HOJA_CARTERA, offset, "aging_no_suma",
                         "suma de buckets=%.2f vs total_cartera=%.2f" % (suma, total),
                         {"suma": suma, "total": total})
            )

        fila = {
            "activo_key": activo,
            "periodo": periodo,
            "fecha_corte_fuente": _fecha_fuente(corte),
            "cliente": _texto(_get(row, cols, "cliente")),
            "marca": _texto(_get(row, cols, "marca")),
            "identificacion": (
                _texto(_get(row, cols, "identificacion"))
                or _texto(_get(row, cols, "identificación"))
            ),
            "documento": _texto(_get(row, cols, "documento")),
            "concepto": _texto(_get(row, cols, "concepto")),
            "inmueble": _texto(_get(row, cols, "inmueble")),
            "fecha_vencimiento": _fecha_fuente(_get(row, cols, "fecha_vencimiento")),
            "dias_vencimiento": (
                _num(_get(row, cols, "dias_de_vencimiento"))
                or _num(_get(row, cols, "días_de_vencimiento"))
            ),
            "saldo_a_favor": _num(_get(row, cols, "saldo_a_favor")),
            "total_cartera": total,
            "source_row": offset,
        }
        fila.update(buckets)
        hoja.filas.append(fila)
    return hoja


def parse_recaudado(rows: list) -> HojaParseada:
    hoja = HojaParseada(HOJA_RECAUDADO)
    if not rows:
        return hoja
    cols = _col_map(rows[0])

    for offset, row in enumerate(rows[1:], start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        hoja.filas_leidas += 1
        activo = _resolver_activo(
            _get(row, cols, "Proyecto"), HOJA_RECAUDADO, offset, hoja.anomalias
        )
        corte = _get(row, cols, "Fecha corte")
        periodo = _periodo(corte)
        if periodo is None:
            hoja.anomalias.append(
                Anomalia(HOJA_RECAUDADO, offset, "sin_fecha_corte",
                         "Fecha corte vacia o ilegible", corte)
            )
        monto = _num(_get(row, cols, "Valor recaudado"))
        if monto is None:
            hoja.anomalias.append(
                Anomalia(HOJA_RECAUDADO, offset, "sin_monto", "Valor recaudado vacio")
            )

        hoja.filas.append({
            "activo_key": activo,
            "periodo": periodo,
            "fecha_corte_fuente": _fecha_fuente(corte),
            "monto_uf": monto,
            "source_row": offset,
        })
    return hoja


# -- entrada publica ---------------------------------------------------------

_PARSERS = {
    HOJA_RENT_ROLL: parse_rent_roll,
    HOJA_AUXILIAR: parse_auxiliar_contable,
    HOJA_CARTERA: parse_cartera,
    HOJA_RECAUDADO: parse_recaudado,
}


def parse_planilla(filepath: str) -> PlanillaParseada:
    """Lee las cuatro hojas en alcance y devuelve filas normalizadas + anomalias."""
    wb = _abrir(filepath)
    try:
        faltantes = [h for h in HOJAS_FIRMA if h not in wb.sheetnames]
        if faltantes:
            raise ValueError(
                "No parece una planilla JLL v2: faltan las hojas %s" % faltantes
            )
        parseadas = {}
        for nombre, parser in _PARSERS.items():
            if nombre not in wb.sheetnames:
                vacia = HojaParseada(nombre)
                vacia.anomalias.append(
                    Anomalia(nombre, 0, "hoja_ausente",
                             "El archivo no trae la hoja %s" % nombre)
                )
                parseadas[nombre] = vacia
                continue
            rows = list(wb[nombre].iter_rows(values_only=True))
            parseadas[nombre] = parser(rows)
    finally:
        wb.close()

    return PlanillaParseada(
        source_file=filepath,
        rent_roll=parseadas[HOJA_RENT_ROLL],
        movimientos=parseadas[HOJA_AUXILIAR],
        cartera=parseadas[HOJA_CARTERA],
        recaudacion=parseadas[HOJA_RECAUDADO],
    )
